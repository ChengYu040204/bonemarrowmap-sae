#!/usr/bin/env python3
"""
Generate Excel comparison: Pearson RAG vs t-test RAG for ANOVA top 10.
Run after task7c_ttest_annotation.py and compare_ttest_results.py.

Usage:
  python3 make_ttest_eval.py

Output:
  Datasheets/20260512_SAE_TTest_vs_Pearson_Eval.xlsx
"""

import json, os
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

os.makedirs("Datasheets", exist_ok=True)
OUTPUT = "Datasheets/20260512_SAE_TTest_vs_Pearson_Eval.xlsx"

# ============================================================
# Load data
# ============================================================
with open("sae_annotations.json") as f:
    ann_z = {a["feature_id"]: a for a in json.load(f)}

with open("sae_annotations_rag.json") as f:
    ann_p = {a["feature_id"]: a for a in json.load(f)}

with open("sae_annotations_rag_ttest.json") as f:
    ann_t_list = json.load(f)
    ann_t = {a["feature_id"]: a for a in ann_t_list}

with open("gene_signatures.json") as f:
    raw_p = json.load(f)
    pearson_sigs = {int(k): v["genes"] for k, v in raw_p.items()}

with open("gene_signatures_ttest.json") as f:
    raw_t = json.load(f)
    ttest_sigs = {int(k): v["genes"] for k, v in raw_t.items()}

df_clusters = pd.read_csv("task6_cluster_assignments.csv")

# ANOVA top 10 feature IDs (by anova_rank in zero-shot annotations)
anova_fids = sorted(ann_z.keys(), key=lambda f: ann_z[f].get("anova_rank") or 999)[:10]

# Human ground truth (Task 3)
HUMAN_GT = {
    799: "Large Pre-B S-phase burst (STMN1/PCNA/IGLL1)",
    73:  "Plasma cell UPR initiation (HSP90B1/HSPA5/DDIT3)",
    75:  "Plasma cell UPR / chaperone (HSP family/XBP1)",
    8:   "Plasma cell ER translocon entry (SEC61/SSR4)",
    56:  "Plasma cell mTORC1 + ER stress coordination",
    24:  "Plasma cell ERAD / protein quality control",
    1:   "Terminal B-cell diff + IgG secretion (IGKC/IGLC)",
    7:   "BM niche survival signalling (BCMA/SLAMF7/CD27)",
    6:   "Plasma cell UPR + PRDX4 ROS scavenging",
    16:  "Plasma cell IgLambda production (IGLC2/IGLC3)",
}

# ============================================================
# Helpers
# ============================================================
def cell(ws, row, col, value="", bold=False, color=None, bg=None,
         align="left", wrap=False, size=10):
    c = ws.cell(row=row, column=col, value=value)
    c.font = Font(bold=bold, size=size,
                  color=color if color else "1A1A2E")
    if bg:
        c.fill = PatternFill("solid", fgColor=bg)
    c.alignment = Alignment(horizontal=align, vertical="top",
                             wrap_text=wrap)
    return c

def thin_border():
    s = Side(style="thin", color="CCCCCC")
    return Border(left=s, right=s, top=s, bottom=s)

def apply_border(ws, min_row, max_row, min_col, max_col):
    for r in range(min_row, max_row + 1):
        for c in range(min_col, max_col + 1):
            ws.cell(r, c).border = thin_border()

BG = {
    "header":  "1B3A5C",
    "subhdr":  "D6E8F7",
    "best":    "D4EFDF",   # green: t-test improvement
    "partial": "FFF9C4",   # yellow: partial match
    "fail":    "FDECEA",   # red: t-test no improvement
    "light":   "F5F7FA",
    "white":   "FFFFFF",
}
FG = {
    "white":  "FFFFFF",
    "dark":   "1B3A5C",
    "green":  "1E8449",
    "coral":  "C0392B",
    "amber":  "D68910",
    "blue":   "1E7EC8",
    "gray":   "6C7A89",
}

# ============================================================
# Sheet 1: Side-by-Side Comparison (ANOVA Top 10)
# ============================================================
wb = Workbook()
ws1 = wb.active
ws1.title = "Top10 Comparison"

# Title
ws1.merge_cells("A1:J1")
cell(ws1, 1, 1, "ANOVA Top 10: Pearson RAG vs t-test RAG Annotation Comparison",
     bold=True, color=FG["white"], bg=BG["header"], align="center", size=12)
ws1.row_dimensions[1].height = 22

# Column headers
headers = [
    "Rank", "Feature", "F-stat",
    "Pearson Top Genes (5)", "t-test Top Genes (5)",
    "Human GT", "RAG Pearson", "RAG t-test", "Improvement?", "Notes"
]
col_widths = [6, 7, 10, 28, 28, 36, 34, 34, 14, 32]
for j, (h, w) in enumerate(zip(headers, col_widths), 1):
    cell(ws1, 2, j, h, bold=True, color=FG["white"],
         bg=BG["subhdr"] if j > 3 else BG["header"],
         align="center", size=9)
    ws1.column_dimensions[get_column_letter(j)].width = w
ws1.row_dimensions[2].height = 18

# Data rows
for rank, fid in enumerate(anova_fids, 1):
    row = rank + 2
    ws1.row_dimensions[row].height = 60

    z_name = ann_z.get(fid, {}).get("programme_name", "")
    p_name = ann_p.get(fid, {}).get("programme_name", "")
    t_name = ann_t.get(fid, {}).get("programme_name", "")
    human  = HUMAN_GT.get(fid, "")
    f_stat = ann_z.get(fid, {}).get("f_stat", "")

    p_genes = ", ".join(pearson_sigs.get(fid, [])[:5])
    t_genes = ", ".join(ttest_sigs.get(fid,   [])[:5])

    # Assess improvement
    h_lower = human.lower()
    p_match = any(w in p_name.lower() for w in ["upr", "erad", "mtorc", "translocon", "pre-b", "lambda", "survival", "ros"])
    t_match = any(w in t_name.lower() for w in ["upr", "erad", "mtorc", "translocon", "pre-b", "lambda", "survival", "ros"])

    if t_match and not p_match:
        improvement = "YES"
        row_bg = BG["best"]
    elif t_match and p_match:
        improvement = "BOTH"
        row_bg = BG["partial"]
    elif not t_match and not p_match:
        improvement = "NEITHER"
        row_bg = BG["fail"]
    else:
        improvement = "PEARSON better"
        row_bg = BG["fail"]

    values = [rank, f"F#{fid}", f_stat, p_genes, t_genes, human, p_name, t_name, improvement, ""]
    for j, v in enumerate(values, 1):
        bg = row_bg if j > 3 else BG["white"]
        c = cell(ws1, row, j, v, bg=bg, align="center" if j <= 3 else "left",
                 wrap=True, size=9)

apply_border(ws1, 1, len(anova_fids) + 2, 1, len(headers))
ws1.freeze_panes = "A3"

# ============================================================
# Sheet 2: Cluster 2 Full Breakdown (all 41 features)
# ============================================================
ws2 = wb.create_sheet("Cluster2 Plasma Cell")
ws2.merge_cells("A1:F1")
cell(ws2, 1, 1, "Cluster 2 (Plasma Cell): All 41 Features — Pearson vs t-test",
     bold=True, color=FG["white"], bg=BG["header"], align="center", size=12)
ws2.row_dimensions[1].height = 22

c2_headers = ["Feature", "Cluster", "Pearson Top Genes (5)", "t-test Top Genes (5)", "RAG Pearson Name", "RAG t-test Name"]
c2_widths   = [8, 8, 35, 35, 38, 38]
for j, (h, w) in enumerate(zip(c2_headers, c2_widths), 1):
    cell(ws2, 2, j, h, bold=True, color=FG["white"], bg=BG["subhdr"], align="center", size=9)
    ws2.column_dimensions[get_column_letter(j)].width = w
ws2.row_dimensions[2].height = 18

c2_fids = df_clusters[df_clusters["cluster"] == 2]["feature_id"].tolist()
for i, fid in enumerate(c2_fids):
    row = i + 3
    ws2.row_dimensions[row].height = 45
    p_name = ann_p.get(fid, {}).get("programme_name", "")
    t_name = ann_t.get(fid, {}).get("programme_name", "")
    p_genes = ", ".join(pearson_sigs.get(fid, [])[:5])
    t_genes = ", ".join(ttest_sigs.get(fid,   [])[:5])

    same_name = (p_name == t_name)
    row_bg = BG["light"] if same_name else BG["best"]

    vals = [f"F#{fid}", 2, p_genes, t_genes, p_name, t_name]
    for j, v in enumerate(vals, 1):
        cell(ws2, row, j, v, bg=row_bg if j > 2 else BG["white"],
             align="left", wrap=True, size=9)

apply_border(ws2, 1, len(c2_fids) + 2, 1, 6)
ws2.freeze_panes = "A3"

# ============================================================
# Sheet 3: Summary Statistics
# ============================================================
ws3 = wb.create_sheet("Summary Statistics")
ws3.column_dimensions["A"].width = 42
ws3.column_dimensions["B"].width = 18
ws3.column_dimensions["C"].width = 18
ws3.column_dimensions["D"].width = 28

ws3.merge_cells("A1:D1")
cell(ws3, 1, 1, "Summary: t-test vs Pearson RAG Annotation Quality",
     bold=True, color=FG["white"], bg=BG["header"], align="center", size=12)
ws3.row_dimensions[1].height = 22

# Section A: Name diversity
cell(ws3, 2, 1, "A. Programme Name Diversity (all 209 features)",
     bold=True, bg=BG["subhdr"], size=10)
ws3.merge_cells("A2:D2")

stats_a = [
    ("Unique names — Zero-shot",    len(set(a["programme_name"] for a in ann_z.values())),  "", ""),
    ("Unique names — RAG Pearson",  len(set(a["programme_name"] for a in ann_p.values())),  "", ""),
    ("Unique names — RAG t-test",   len(set(a["programme_name"] for a in ann_t.values())),  "", ""),
]
for i, (label, val, _, note) in enumerate(stats_a, 3):
    cell(ws3, i, 1, label, size=9)
    cell(ws3, i, 2, val, align="center", size=9,
         bold=(i == 5),
         color=FG["green"] if i == 5 else FG["dark"])

# Section B: Cluster 2 specific
cell(ws3, 7, 1, "B. Cluster 2 (Plasma Cell, 41 features) — Sub-programme Resolution",
     bold=True, bg=BG["subhdr"], size=10)
ws3.merge_cells("A7:D7")

c2_p_names = [ann_p.get(f, {}).get("programme_name", "") for f in c2_fids]
c2_t_names = [ann_t.get(f, {}).get("programme_name", "") for f in c2_fids]

def mean_jacc(sigs, fids):
    vals = []
    fids = [f for f in fids if f in sigs]
    for i in range(len(fids)):
        for j in range(i + 1, len(fids)):
            a = set(sigs[fids[i]]) if isinstance(sigs[fids[i]], list) else sigs[fids[i]]
            b = set(sigs[fids[j]]) if isinstance(sigs[fids[j]], list) else sigs[fids[j]]
            if a or b:
                vals.append(len(a & b) / len(a | b))
    return round(np.mean(vals), 3) if vals else 0

p_jac = mean_jacc(pearson_sigs, c2_fids)
t_jac = mean_jacc(ttest_sigs,   c2_fids)

stats_b = [
    ("Unique programme names — Pearson",   len(set(c2_p_names)), "", ""),
    ("Unique programme names — t-test",    len(set(c2_t_names)), "", "Higher = better sub-programme distinction"),
    ("Mean pairwise Jaccard — Pearson",    p_jac,                "", "Gene list similarity across features"),
    ("Mean pairwise Jaccard — t-test",     t_jac,                "", "Lower = more feature-specific gene lists"),
]
for i, (label, val, _, note) in enumerate(stats_b, 8):
    cell(ws3, i, 1, label, size=9)
    cell(ws3, i, 2, val,   align="center", size=9)
    cell(ws3, i, 4, note,  size=8, color=FG["gray"])

# Section C: Interpretation
cell(ws3, 13, 1, "C. Interpretation",
     bold=True, bg=BG["subhdr"], size=10)
ws3.merge_cells("A13:D13")
interp = (
    "t-test identifies genes specifically upregulated in the top-200 activating cells "
    "vs a random background of 2,000 cells, rather than computing global Pearson correlation "
    "across all 161,604 cells. This exposes within-cluster differences (e.g. UPR vs ERAD vs "
    "mTORC1 within plasma cells) that are invisible to global correlation because all plasma "
    "cell features share the same global activation pattern."
)
c = ws3.cell(14, 1, interp)
c.font = Font(size=9, color=FG["dark"])
c.alignment = Alignment(wrap_text=True, vertical="top")
c.fill = PatternFill("solid", fgColor=BG["light"])
ws3.merge_cells("A14:D14")
ws3.row_dimensions[14].height = 60

apply_border(ws3, 1, 14, 1, 4)

wb.save(OUTPUT)
print(f"Saved: {OUTPUT}")
